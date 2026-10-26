#!/usr/bin/env python3
"""build_sheet.py 결과를 LibreOffice(headless)로 재계산해 brief 부록 C 케이스를 원 단위로 검증.

사용:  python3 sheet/verify_sheet.py [--xlsx sheet/out/card-deduction-tracker-2027.xlsx]
필요:  openpyxl, LibreOffice(`soffice`)
방식:  openpyxl로 입력 칸에 값을 넣어 저장 → `soffice --headless --convert-to xlsx`(로드 시 수식 재계산)
       → openpyxl data_only로 결과 셀을 읽어 기대값과 비교.

검증 항목
  1) '2026마무리' C15~C23: 부록 C 케이스 1·2·3·4 + 케이스 3 K=0, 케이스 4 G=70,000,001 (원 단위, 결과 문구 글자 단위)
  2) '설정' C4 토글 ON 이어도 '2026마무리' 값 불변
  3) '2027공제' 토글 OFF: 가계부에 케이스 1 금액(신용 16,000,000 + 체크 3,000,000), 총급여 50,000,000 → 1,425,000
  4) '2027공제' 토글 OFF/ON: 전통시장·대중교통·문화체육 섞인 가계부를 독립 구현(파이썬 Fraction)과 대조
  5) 문화체육 칸(B-4)이 있는 '2026마무리' 계산을 독립 구현과 대조
  6) 전 시트 수식 오류(#REF!, #NAME?, #VALUE!, #DIV/0!, Err:xxx) 0개
"""
from __future__ import annotations

import argparse
import datetime as dt
import shutil
import subprocess
import sys
import tempfile
from fractions import Fraction
from math import floor
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
TOGGLE_ON = "☑ 켬 — 정부 개정안 시나리오"

# ---------------- 독립 구현 (brief B-2 / B-4, 설정 탭 변수) ----------------
CUR = dict(min_rate=25, rc=15, rd=30, rm=40, transit_sep=1, rcu=30, culture_cap=70_000_000, split=70_000_000,
           basic_low=3_000_000, basic_high=2_500_000, child_low=500_000, child_high=250_000, child_max=2,
           extra_low=3_000_000, extra_high=2_000_000)
PROP = dict(CUR, transit_sep=0, culture_cap=0, extra_low=2_000_000, extra_high=1_000_000)


def ref_deduction(G, K, c, d, s, t, p):
    """c 신용, d 체크, s 문화체육(이미 상한 적용), t 전통시장(+대중교통 별도 시)."""
    M = Fraction(G * p["min_rate"], 100)
    total = c + d + s + t
    if total <= M:
        base = Fraction(0)
    else:
        gross = Fraction(c * p["rc"] + d * p["rd"] + s * p["rcu"] + t * p["rm"], 100)
        if c >= M:
            sub = M * p["rc"] / 100
        elif c + d + s >= M:
            sub = Fraction(c * p["rc"], 100) + (M - c) * p["rd"] / 100
        else:
            sub = Fraction(c * p["rc"], 100) + Fraction((d + s) * p["rd"], 100) + (M - c - d - s) * p["rm"] / 100
        base = gross - sub
    k = min(K, p["child_max"])
    low = G <= p["split"]
    BL = (p["basic_low"] + k * p["child_low"]) if low else (p["basic_high"] + k * p["child_high"])
    EL = p["extra_low"] if low else p["extra_high"]
    if base <= BL:
        return floor(base)
    target = Fraction(t * p["rm"], 100) + (Fraction(s * p["rcu"], 100) if low else 0)
    return floor(BL + min(base - BL, target, EL))


def ref_2026(G, C, D, TM, TT, S, K, X):
    p = CUR
    s = S if G <= p["culture_cap"] else 0
    t = TM + TT
    A = ref_deduction(G, K, C + X, D, s, t, p)
    B = ref_deduction(G, K, C, D + X, s, t, p)
    return A, B


def ref_2027(G, K, credit, debit, market, tr_credit, tr_debit, culture, p):
    if p["transit_sep"] == 1:
        c, d, t = credit, debit, market + tr_credit + tr_debit
    else:
        c, d, t = credit + tr_credit, debit + tr_debit, market
    s = culture if (p["culture_cap"] == 0 or G <= p["culture_cap"]) else 0
    return ref_deduction(G, K, c, d, s, t, p)


# ---------------- LibreOffice 재계산 ----------------
class Recalc:
    def __init__(self):
        self.soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not self.soffice:
            raise SystemExit("미검증: LibreOffice(soffice)가 없습니다.")
        self.tmp = Path(tempfile.mkdtemp(prefix="sheetverify-"))
        self.profile = self.tmp / "profile"
        self.n = 0

    def run(self, wb) -> "Workbook":
        self.n += 1
        src = self.tmp / f"in{self.n}.xlsx"
        outdir = self.tmp / f"out{self.n}"
        wb.save(src)
        cmd = [self.soffice, f"-env:UserInstallation=file://{self.profile}", "--headless", "--calc",
               "--convert-to", "xlsx", "--outdir", str(outdir), str(src)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        out = outdir / src.name
        if not out.exists():
            raise RuntimeError(f"LibreOffice 변환 실패: {r.stdout}\n{r.stderr}")
        return load_workbook(out, data_only=True)

    def cleanup(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


ERR_PREFIX = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "#ERROR", "Err:")


def scan_errors(wb) -> list[str]:
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERR_PREFIX):
                    bad.append(f"{ws.title}!{c.coordinate}={c.value}")
    return bad


def won(n):
    return f"{n:,}"


CASES = {
    "1": dict(G=50_000_000, C=12_000_000, D=3_000_000, TM=0, K=0, X=4_000_000,
              exp=dict(C15=12_500_000, C16=0, C17=0, C18=3_000_000, C19=3_000_000, C20=1_425_000, C21=1_950_000, C22=525_000),
              msg="남은 4,000,000원을 체크카드·현금영수증으로 쓰면 소득공제가 525,000원 늘어납니다 (신용 1,425,000원 → 체크 1,950,000원)."),
    "2": dict(G=40_000_000, C=6_000_000, D=1_000_000, TM=500_000, K=0, X=3_000_000,
              exp=dict(C15=10_000_000, C16=2_500_000, C17=0, C20=200_000, C21=200_000, C22=0),
              msg="남은 지출은 대부분 25% 문턱을 채우는 데 쓰여 결제수단에 따른 차이가 없습니다."),
    "3": dict(G=80_000_000, C=30_000_000, D=10_000_000, TM=2_000_000, K=1, X=5_000_000,
              exp=dict(C15=20_000_000, C16=0, C18=2_750_000, C19=2_000_000, C20=3_550_000, C21=3_550_000, C22=0),
              msg="이미 공제 한도에 도달해 체크카드로 바꿔도 공제액이 같습니다. 혜택 좋은 신용카드를 쓰세요."),
    "3 (K=0)": dict(G=80_000_000, C=30_000_000, D=10_000_000, TM=2_000_000, K=0, X=5_000_000,
                    exp=dict(C18=2_500_000, C20=3_300_000, C21=3_300_000, C22=0),
                    msg="이미 공제 한도에 도달해 체크카드로 바꿔도 공제액이 같습니다. 혜택 좋은 신용카드를 쓰세요."),
    "4": dict(G=70_000_000, C=40_000_000, D=0, TM=0, K=0, X=0,
              exp=dict(C15=17_500_000, C18=3_000_000, C20=3_000_000, C21=3_000_000, C22=0),
              msg="이미 공제 한도에 도달해 체크카드로 바꿔도 공제액이 같습니다. 혜택 좋은 신용카드를 쓰세요."),
    "4 (G=70,000,001)": dict(G=70_000_001, C=40_000_000, D=0, TM=0, K=0, X=0,
                             exp=dict(C18=2_500_000, C20=2_500_000, C21=2_500_000, C22=0),
                             msg="이미 공제 한도에 도달해 체크카드로 바꿔도 공제액이 같습니다. 혜택 좋은 신용카드를 쓰세요."),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(HERE / "out" / "card-deduction-tracker-2027.xlsx"))
    a = ap.parse_args()
    xlsx = Path(a.xlsx)
    if not xlsx.exists():
        subprocess.run([sys.executable, str(HERE / "build_sheet.py"), "--out", str(xlsx)], check=True)

    rc = Recalc()
    results: list[tuple[bool, str]] = []

    def check(name, ok, detail=""):
        results.append((ok, f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  — {detail}" if detail else "")))

    def fresh():
        return load_workbook(xlsx)

    def set_2026(wb, c, toggle=False, S=0, TT=0):
        ws = wb["2026마무리"]
        ws["C5"], ws["C6"], ws["C7"], ws["C8"], ws["C9"] = c["G"], c["C"], c["D"], c["TM"], TT
        ws["C10"], ws["C11"], ws["C12"] = S, c["K"], c["X"]
        if toggle:
            wb["설정"]["C4"] = TOGGLE_ON

    try:
        # 0) 빈 시트 오류 스캔
        out = rc.run(fresh())
        errs = scan_errors(out)
        check("빈 시트 수식 오류 0개", not errs, ", ".join(errs[:10]))
        check("빈 시트 2026마무리 C23 안내", out["2026마무리"]["C23"].value == "노란 칸을 채우면 결과가 나옵니다.", str(out["2026마무리"]["C23"].value))

        # 1) 부록 C 케이스
        for name, c in CASES.items():
            wb = fresh()
            set_2026(wb, c)
            out = rc.run(wb)
            ws = out["2026마무리"]
            got = {k: ws[k].value for k in c["exp"]}
            ok_vals = all(got[k] == v for k, v in c["exp"].items())
            A, B = ref_2026(c["G"], c["C"], c["D"], c["TM"], 0, 0, c["K"], c["X"])
            ok_ref = ws["C20"].value == A and ws["C21"].value == B
            check(f"케이스 {name} 금액(원 단위)", ok_vals and ok_ref, "C20/C21/C22 = " + " / ".join(won(ws[k].value) for k in ("C20", "C21", "C22")))
            check(f"케이스 {name} 결과 문구", ws["C23"].value == c["msg"], str(ws["C23"].value))
            errs = scan_errors(out)
            check(f"케이스 {name} 수식 오류 0개", not errs, ", ".join(errs[:5]))

        # 2) 토글 ON이어도 2026마무리 불변
        wb = fresh()
        set_2026(wb, CASES["1"], toggle=True)
        out = rc.run(wb)
        ws = out["2026마무리"]
        check("토글 ON → 2026마무리 불변(케이스 1)", (ws["C20"].value, ws["C21"].value, ws["C22"].value) == (1_425_000, 1_950_000, 525_000))
        check("토글 ON → 설정 C5 = 1", out["설정"]["C5"].value == 1)

        # 5) 문화체육·대중교통 칸이 있는 2026마무리 (B-4)
        for G, C_, D_, X_ in ((50_000_000, 9_000_000, 2_000_000, 4_000_000), (80_000_000, 20_000_000, 4_000_000, 3_000_000)):
            c = dict(G=G, C=C_, D=D_, TM=700_000, K=0, X=X_)
            wb = fresh()
            set_2026(wb, c, S=1_200_000, TT=900_000)
            out = rc.run(wb)
            ws = out["2026마무리"]
            A, B = ref_2026(G, C_, D_, 700_000, 900_000, 1_200_000, 0, X_)
            check(f"2026마무리 문화체육 포함(G={won(G)}) 독립 구현 대조", (ws["C20"].value, ws["C21"].value) == (A, B),
                  f"시트 {won(ws['C20'].value)}/{won(ws['C21'].value)} vs 기대 {won(A)}/{won(B)}")
            if G > 70_000_000:
                check("7천만원 초과 + 문화비 입력 → 경고 문구", str(ws["C25"].value).startswith("총급여 7천만원 초과"))

        # 3) 2027공제 토글 OFF, 케이스 1 금액
        def ledger(wb, rows):
            ws = wb["2027가계부"]
            for i, (m, kind, method, amt) in enumerate(rows):
                r = 5 + i
                ws[f"A{r}"] = dt.date(2027, m, 15)
                ws[f"C{r}"] = kind
                ws[f"D{r}"] = "기타"
                ws[f"E{r}"] = method
                ws[f"F{r}"] = amt

        wb = fresh()
        ledger(wb, [(1, "지출", "신용카드", 16_000_000), (2, "지출", "체크·현금영수증", 3_000_000), (2, "수입", "", 4_000_000)])
        wb["2027공제"]["C5"] = 50_000_000
        out = rc.run(wb)
        d = out["2027공제"]
        check("2027공제 OFF 케이스 1 → 1,425,000", d["C21"].value == 1_425_000, won(d["C21"].value))
        m = out["2027월별"]
        check("2027월별 1월 신용 16,000,000 / 2월 체크 3,000,000 / 수입 집계",
              m["D5"].value == 16_000_000 and m["E6"].value == 3_000_000 and m["B6"].value == 4_000_000 and m["L17"].value == 19_000_000,
              f"D5={m['D5'].value} E6={m['E6'].value} B6={m['B6'].value} L17={m['L17'].value}")
        check("2027월별 진행률(19,000,000/12,500,000 → 100%)", m["M17"].value == 1, str(m["M17"].value))
        check("2027공제 C26 = 현행 규칙", d["C26"].value == "현행 규칙", str(d["C26"].value))

        # 4) 2027공제 OFF/ON 혼합 가계부
        rows = [(1, "지출", "신용카드", 18_000_000), (3, "지출", "체크·현금영수증", 6_000_000), (4, "지출", "전통시장", 1_500_000),
                (5, "지출", "대중교통(신용)", 800_000), (6, "지출", "대중교통(체크)", 700_000), (7, "지출", "문화체육", 1_000_000),
                (8, "지출", "공제 제외", 2_000_000), (9, "지출", "현금(영수증 없음)", 300_000)]
        for G, K in ((60_000_000, 1), (90_000_000, 2)):
            for on in (False, True):
                wb = fresh()
                ledger(wb, rows)
                wb["2027공제"]["C5"] = G
                wb["2027공제"]["C6"] = K
                if on:
                    wb["설정"]["C4"] = TOGGLE_ON
                out = rc.run(wb)
                d = out["2027공제"]
                exp = ref_2027(G, K, 18_000_000, 6_000_000, 1_500_000, 800_000, 700_000, 1_000_000, PROP if on else CUR)
                check(f"2027공제 G={won(G)} K={K} 토글 {'ON' if on else 'OFF'} 독립 구현 대조", d["C21"].value == exp,
                      f"시트 {won(d['C21'].value)} vs 기대 {won(exp)}")
                label_ok = (d["C26"].value == "정부 개정안(국회 미통과) — 미확정 시나리오") if on else (d["C26"].value == "현행 규칙")
                check(f"2027공제 적용 규칙 표기 ({'ON' if on else 'OFF'})", label_ok, str(d["C26"].value))
                errs = scan_errors(out)
                check(f"2027공제 G={won(G)} {'ON' if on else 'OFF'} 수식 오류 0개", not errs, ", ".join(errs[:5]))
    finally:
        rc.cleanup()

    for _, line in results:
        print(line)
    failed = sum(1 for ok, _ in results if not ok)
    print(f"\n{'ALL PASS' if not failed else str(failed) + ' FAILED'} ({len(results)} checks, LibreOffice 재계산 {rc.n}회)")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
